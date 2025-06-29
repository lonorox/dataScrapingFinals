"""
Reports generation module for scraped data.
Provides automated reports with insights and data export in multiple formats.
"""

import pandas as pd
import numpy as np
import sqlite3
import json
import base64
from datetime import datetime, timedelta
from typing import Dict, List, Any, Optional, Tuple
import os
from pathlib import Path
import matplotlib.pyplot as plt
import seaborn as sns
from collections import Counter
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
import jinja2
import webbrowser

# Import analysis modules
from src.analysis.statistic import DataStatistics
from src.analysis.trends import TrendAnalysis

# Import constants and templates
from src.analysis.constants import HTML_REPORT_TEMPLATE, VISUALIZATION_IMAGES, ALERT_STYLES, REPORT_SETTINGS


class ReportGenerator:
    """Generate comprehensive reports with insights and data export."""
    
    def __init__(self, data_path: str = "data_output/raw",
                 db_path: str = "data_output/raw/scraped_articles.db"):
        """Initialize with data paths."""
        self.data_path = data_path
        self.db_path = db_path
        self.df = None
        self.connection = None
        
        # Initialize analysis modules
        self.statistics = DataStatistics(os.path.join(data_path, "combined.csv"), db_path)
        self.trends = TrendAnalysis(os.path.join(data_path, "combined.csv"), db_path)
        
    def load_data(self) -> bool:
        """Load data from CSV or database."""
        try:
            # Check for combined.csv in data_output directory
            combined_csv_path = os.path.join(self.data_path, "combined.csv")
            if os.path.exists(combined_csv_path):
                self.df = pd.read_csv(combined_csv_path)
                print(f"Loaded {len(self.df)} records from combined.csv")
            elif os.path.exists(self.db_path):
                self.connection = sqlite3.connect(self.db_path)
                self.df = pd.read_sql_query("SELECT * FROM articles", self.connection)
                print(f"Loaded {len(self.df)} records from database")
            else:
                print(f"No data files found in {self.data_path}")
                print(f"Available files: {[f for f in os.listdir(self.data_path) if os.path.isfile(os.path.join(self.data_path, f))]}")
                return False
            
            # Preprocess dates
            if 'publication_date_datetime' in self.df.columns:
                self.df['pub_date'] = pd.to_datetime(self.df['publication_date_datetime'], errors='coerce')
            
            # Load data into analysis modules
            self.statistics.load_data()
            self.trends.load_data()
            
            return True
        except Exception as e:
            print(f"Error loading data: {e}")
            return False
    
    def close_connection(self):
        """Close database connection."""
        if self.connection:
            self.connection.close()
        self.statistics.close_connection()
        self.trends.close_connection()
    
    def load_visualization_images(self, reports_dir: str = "data_output/processed") -> Dict[str, str]:
        """Load and encode visualization images as base64 strings."""
        images = {}
        
        if not os.path.exists(reports_dir):
            print(f"Reports directory not found: {reports_dir}")
            return images
        
        for image_key, filename in VISUALIZATION_IMAGES.items():
            image_path = os.path.join(reports_dir, filename)
            if os.path.exists(image_path):
                try:
                    with open(image_path, 'rb') as img_file:
                        img_data = img_file.read()
                        img_base64 = base64.b64encode(img_data).decode('utf-8')
                        images[image_key] = img_base64
                        print(f"Loaded visualization: {filename}")
                except Exception as e:
                    print(f"Error loading image {filename}: {e}")
            else:
                print(f"Image not found: {image_path}")
        
        return images
    
    def generate_data_quality_issues(self) -> List[str]:
        """Generate list of data quality issues found in the dataset."""
        issues = []
        
        if self.df is None:
            return ["No data loaded"]
        
        # Get data quality report from statistics module
        quality_report = self.statistics.data_quality_check()
        
        # Check for missing titles
        if "title" in quality_report.get("missing_values", {}):
            missing_titles = quality_report["missing_values"]["title"]["count"]
            if missing_titles > 0:
                issues.append(f"{missing_titles} articles missing titles")
        
        # Check for missing summaries
        if "summary" in quality_report.get("missing_values", {}):
            missing_summaries = quality_report["missing_values"]["summary"]["count"]
            if missing_summaries > 0:
                issues.append(f"{missing_summaries} articles missing summaries")
        
        # Check for missing authors
        if "author" in quality_report.get("missing_values", {}):
            missing_authors = quality_report["missing_values"]["author"]["count"]
            if missing_authors > 0:
                issues.append(f"{missing_authors} articles missing authors")
        
        # Check for duplicates
        if "duplicates" in quality_report:
            duplicates = quality_report["duplicates"]["total"]
            if duplicates > 0:
                issues.append(f"{duplicates} duplicate articles found")
            
            url_duplicates = quality_report["duplicates"].get("url_duplicates", 0)
            if url_duplicates > 0:
                issues.append(f"{url_duplicates} duplicate URLs found")
        
        # Check for very short titles
        if 'title' in self.df.columns:
            short_titles = (self.df['title'].str.len() < 10).sum()
            if short_titles > 0:
                issues.append(f"{short_titles} articles have very short titles (< 10 characters)")
        
        return issues
    
    def generate_executive_summary(self) -> Dict[str, Any]:
        """Generate executive summary with key insights."""
        if self.df is None:
            return {"error": "No data loaded"}
        
        # Get statistical summary from statistics module
        stats_summary = self.statistics.statistical_summary()
        
        summary = {
            "overview": {},
            "key_insights": [],
            "recommendations": [],
            "performance_metrics": {}
        }
        
        # Overview
        summary["overview"] = {
            "total_articles": len(self.df),
            "date_range": {
                "start": str(self.df['pub_date'].min()) if 'pub_date' in self.df.columns else "Unknown",
                "end": str(self.df['pub_date'].max()) if 'pub_date' in self.df.columns else "Unknown"
            },
            "unique_sources": self.df['source'].nunique(),
            "unique_authors": self.df['author'].nunique() if 'author' in self.df.columns else 0,
            "source_types": self.df['source_type'].value_counts().to_dict()
        }
        
        # Key insights
        insights = []
        
        # Insight 1: Top performing source
        top_source = self.df['source'].value_counts().index[0]
        top_source_count = self.df['source'].value_counts().iloc[0]
        insights.append(f"Top performing source: {top_source} with {top_source_count} articles")
        
        # Insight 2: Content volume trend
        if 'pub_date' in self.df.columns:
            daily_counts = self.df.groupby(self.df['pub_date'].dt.date).size()
            avg_daily = daily_counts.mean()
            insights.append(f"Average daily publication rate: {avg_daily:.1f} articles")
            
            if len(daily_counts) > 1:
                trend = "increasing" if daily_counts.iloc[-1] > daily_counts.iloc[0] else "decreasing"
                insights.append(f"Publication trend: {trend}")
        
        # Insight 3: Content quality
        if 'title' in self.df.columns:
            avg_title_length = self.df['title'].str.len().mean()
            insights.append(f"Average title length: {avg_title_length:.1f} characters")
        
        # Insight 4: Source diversity
        source_concentration = (self.df['source'].value_counts().head(5).sum() / len(self.df)) * 100
        insights.append(f"Top 5 sources account for {source_concentration:.1f}% of all content")
        
        summary["key_insights"] = insights
        
        # Recommendations
        recommendations = []
    
        summary["recommendations"] = recommendations
        
        # Performance metrics
        summary["performance_metrics"] = {
            "data_completeness": {
                "titles": float((1 - self.df['title'].isnull().sum() / len(self.df)) * 100),
                "summaries": float((1 - self.df['summary'].isnull().sum() / len(self.df)) * 100) if 'summary' in self.df.columns else 0,
                "authors": float((1 - self.df['author'].isnull().sum() / len(self.df)) * 100) if 'author' in self.df.columns else 0
            },
            "source_efficiency": {
                "articles_per_source": float(len(self.df) / self.df['source'].nunique()),
                "top_source_contribution": float((self.df['source'].value_counts().iloc[0] / len(self.df)) * 100)
            }
        }
        
        return summary
    
    def generate_detailed_analysis(self) -> Dict[str, Any]:
        """Generate detailed analysis report using statistics and trends modules."""
        if self.df is None:
            return {"error": "No data loaded"}
        
        # Get analysis from specialized modules
        stats_summary = self.statistics.statistical_summary()
        distributions = self.statistics.distribution_analysis()
        temporal_trends = self.trends.temporal_trend_analysis()
        source_comparison = self.trends.source_comparative_analysis()
        source_type_analysis = self.trends.source_type_analysis()
        
        analysis = {
            "source_analysis": {
                "top_sources": stats_summary.get("source_analysis", {}).get("top_sources", {}),
                "source_type_distribution": stats_summary.get("source_analysis", {}).get("source_type_distribution", {}),
                "source_performance_metrics": {
                    "avg_articles_per_source": float(len(self.df) / self.df['source'].nunique()),
                    "source_concentration_index": float((self.df['source'].value_counts().head(5).sum() / len(self.df)) * 100)
                }
            },
            "content_analysis": stats_summary.get("content_analysis", {}),
            "temporal_analysis": temporal_trends,
            "quality_analysis": self.statistics.data_quality_check()
        }
        
        return analysis
    
    def export_to_excel(self, output_path: str = "data_output/reports/comprehensive_report.xlsx") -> str:
        """Export comprehensive report to Excel with multiple sheets and charts."""
        if self.df is None:
            return "No data to export"
        
        # Create output directory
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Create Excel writer
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            
            # Sheet 1: Executive Summary
            summary_data = self.generate_executive_summary()
            
            # Create summary DataFrame
            summary_rows = []
            for insight in summary_data.get("key_insights", []):
                summary_rows.append({"Insight": insight})
            
            for rec in summary_data.get("recommendations", []):
                summary_rows.append({"Recommendation": rec})
            
            summary_df = pd.DataFrame(summary_rows)
            summary_df.to_excel(writer, sheet_name='Executive Summary', index=False)
            
            # Sheet 2: Source Analysis
            source_analysis = self.df.groupby('source').agg({
                'title': 'count',
                'source_type': 'first'
            }).rename(columns={'title': 'Article Count'}).sort_values('Article Count', ascending=False)
            
            source_analysis.to_excel(writer, sheet_name='Source Analysis')
            
            # Sheet 3: Temporal Analysis
            if 'pub_date' in self.df.columns:
                daily_counts = self.df.groupby(self.df['pub_date'].dt.date).size().reset_index()
                daily_counts.columns = ['Date', 'Article Count']
                daily_counts.to_excel(writer, sheet_name='Daily Trends', index=False)
            
            # Sheet 4: Content Analysis
            content_data = []
            if 'title' in self.df.columns:
                title_lengths = self.df['title'].str.len()
                content_data.append({
                    'Metric': 'Title Length',
                    'Average': title_lengths.mean(),
                    'Median': title_lengths.median(),
                    'Min': title_lengths.min(),
                    'Max': title_lengths.max()
                })
            
            if 'summary' in self.df.columns:
                summary_lengths = self.df['summary'].str.len()
                content_data.append({
                    'Metric': 'Summary Length',
                    'Average': summary_lengths.mean(),
                    'Median': summary_lengths.median(),
                    'Min': summary_lengths.min(),
                    'Max': summary_lengths.max()
                })
            
            content_df = pd.DataFrame(content_data)
            content_df.to_excel(writer, sheet_name='Content Analysis', index=False)
            
            # Sheet 5: Quality Metrics
            quality_data = []
            quality_data.append({
                'Metric': 'Total Articles',
                'Value': len(self.df)
            })
            quality_data.append({
                'Metric': 'Unique Sources',
                'Value': self.df['source'].nunique()
            })
            quality_data.append({
                'Metric': 'Missing Titles',
                'Value': self.df['title'].isnull().sum()
            })
            if 'summary' in self.df.columns:
                quality_data.append({
                    'Metric': 'Missing Summaries',
                    'Value': self.df['summary'].isnull().sum()
                })
            
            quality_df = pd.DataFrame(quality_data)
            quality_df.to_excel(writer, sheet_name='Quality Metrics', index=False)
            
            # Add charts to the workbook
            workbook = writer.book
            
            # Chart 1: Top Sources Bar Chart
            if len(source_analysis) > 0:
                chart_sheet = workbook.create_sheet('Charts')
                
                # Source distribution chart
                chart1 = BarChart()
                chart1.title = "Top 10 Sources by Article Count"
                chart1.x_axis.title = "Source"
                chart1.y_axis.title = "Article Count"
                
                data = Reference(writer.sheets['Source Analysis'], min_col=2, min_row=1, max_row=11, max_col=2)
                cats = Reference(writer.sheets['Source Analysis'], min_col=1, min_row=2, max_row=11)
                chart1.add_data(data, titles_from_data=True)
                chart1.set_categories(cats)
                
                chart_sheet.add_chart(chart1, "A1")
        
        return output_path
    
    def export_to_csv(self, output_dir: str = "data_output/reports") -> Dict[str, str]:
        """Export data to multiple CSV files."""
        if self.df is None:
            return {"error": "No data to export"}
        
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        exported_files = {}
        
        # Main data export
        main_csv = f"{output_dir}/complete_dataset_{timestamp}.csv"
        self.df.to_csv(main_csv, index=False)
        exported_files["complete_dataset"] = main_csv
        
        # Source summary
        source_summary = self.df.groupby('source').agg({
            'title': 'count',
            'source_type': 'first'
        }).rename(columns={'title': 'article_count'}).reset_index()
        source_csv = f"{output_dir}/source_summary_{timestamp}.csv"
        source_summary.to_csv(source_csv, index=False)
        exported_files["source_summary"] = source_csv
        
        # Daily trends
        if 'pub_date' in self.df.columns:
            daily_trends = self.df.groupby(self.df['pub_date'].dt.date).size().reset_index()
            daily_trends.columns = ['date', 'article_count']
            daily_csv = f"{output_dir}/daily_trends_{timestamp}.csv"
            daily_trends.to_csv(daily_csv, index=False)
            exported_files["daily_trends"] = daily_csv
        
        # Content analysis
        content_analysis = []
        if 'title' in self.df.columns:
            title_stats = self.df['title'].str.len().describe()
            content_analysis.append({
                'metric': 'title_length',
                'count': title_stats['count'],
                'mean': title_stats['mean'],
                'std': title_stats['std'],
                'min': title_stats['min'],
                '25%': title_stats['25%'],
                '50%': title_stats['50%'],
                '75%': title_stats['75%'],
                'max': title_stats['max']
            })
        
        if content_analysis:
            content_df = pd.DataFrame(content_analysis)
            content_csv = f"{output_dir}/content_analysis_{timestamp}.csv"
            content_df.to_csv(content_csv, index=False)
            exported_files["content_analysis"] = content_csv
        
        return exported_files
    
    def generate_html_report(self, output_path: str = "data_output/reports/comprehensive_report.html") -> str:
        """Generate comprehensive HTML report with visualizations."""
        if self.df is None:
            return "No data to generate report"
        
        # Create output directory
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Generate analysis data
        executive_summary = self.generate_executive_summary()
        detailed_analysis = self.generate_detailed_analysis()
        data_quality_issues = self.generate_data_quality_issues()
        
        # Load visualization images
        print("Loading visualization images...")
        visualization_images = self.load_visualization_images()
        
        # Prepare top_sources data in the format expected by the template
        top_sources_data = {}
        source_counts = self.df['source'].value_counts().head(REPORT_SETTINGS['max_top_sources'])
        source_types = self.df.groupby('source')['source_type'].first()
        
        for source in source_counts.index:
            top_sources_data[source] = {
                'article_count': int(source_counts[source]),
                'source_type': str(source_types.get(source, 'Unknown'))
            }
        
        # Prepare template data
        template_data = {
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "overview": executive_summary["overview"],
            "key_insights": executive_summary["key_insights"],
            "recommendations": executive_summary["recommendations"],
            "performance_metrics": executive_summary["performance_metrics"],
            "source_types": executive_summary["overview"]["source_types"],
            "top_sources": top_sources_data,
            "data_quality_issues": data_quality_issues,
            # Add visualization images
            **visualization_images
        }
        
        # Render template
        template = jinja2.Template(HTML_REPORT_TEMPLATE)
        html_content = template.render(**template_data)
        
        # Write to file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        print(f"HTML report generated: {output_path}")
        return output_path
    
    def generate_comprehensive_report(self, output_dir: str = "data_output/reports", processed_dir: str = "data_output/processed") -> Dict[str, str]:
        """Generate comprehensive report in multiple formats with visualizations."""
        if self.df is None:
            return {"error": "No data loaded"}
        
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(processed_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        generated_files = {}
        
        # Generate all reports
        print("Generating executive summary...")
        executive_summary = self.generate_executive_summary()
        
        print("Generating detailed analysis...")
        detailed_analysis = self.generate_detailed_analysis()
        
        print("Checking data quality...")
        data_quality_issues = self.generate_data_quality_issues()
        if data_quality_issues:
            print(f"Found {len(data_quality_issues)} data quality issues")
        
        print("Exporting to Excel...")
        excel_path = self.export_to_excel(f"{processed_dir}/comprehensive_report_{timestamp}.xlsx")
        generated_files["excel"] = excel_path
        
        print("Exporting to CSV...")
        csv_files = self.export_to_csv(processed_dir)
        generated_files["csv"] = csv_files
        
        print("Generating statistical visualizations...")
        self.statistics.generate_visualizations(processed_dir)
        
        print("Generating enhanced HTML report with visualizations...")
        html_path = self.generate_html_report(f"{output_dir}/comprehensive_report_{timestamp}.html")
        generated_files["html"] = html_path
        
        # Save JSON report
        json_report = {
            "generated_at": datetime.now().isoformat(),
            "executive_summary": executive_summary,
            "detailed_analysis": detailed_analysis,
            "data_quality_issues": data_quality_issues,
            "visualization_images_loaded": list(self.load_visualization_images().keys())
        }
        
        json_path = f"{processed_dir}/comprehensive_report_{timestamp}.json"
        with open(json_path, 'w') as f:
            json.dump(json_report, f, indent=2, default=str)
        generated_files["json"] = json_path
        
        return generated_files


def main():
    """Main function for running report generation."""
    reporter = ReportGenerator()
    
    if reporter.load_data():
        print("Generating comprehensive reports...")
        
        # Generate all reports
        generated_files = reporter.generate_comprehensive_report()
        
        print("Report generation completed!")
        print("Generated files:")
        for file_type, path in generated_files.items():
            if isinstance(path, dict):
                print(f"  {file_type.upper()}:")
                for sub_type, sub_path in path.items():
                    print(f"    {sub_type}: {sub_path}")
            else:
                print(f"  {file_type.upper()}: {path}")
        
        reporter.close_connection()
    else:
        print("Failed to load data for report generation")


if __name__ == "__main__":
    main() 